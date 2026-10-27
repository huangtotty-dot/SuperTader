# -*- coding: utf-8 -*-
from pathlib import Path
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"E:\superTrader\data_cleaning")
OUT = ROOT / "网络安全概念深度清洗_20261027.xlsx"
audit = json.loads((ROOT / "cybersecurity_audit_20261027.json").read_text(encoding="utf-8"))
current = {x["code"]: x for x in json.loads((ROOT / "cybersecurity_current_20261027.json").read_text(encoding="utf-8"))}


def pair_text(pairs, cyber_only=False):
    items = [p for p in pairs if not cyber_only or p["category"] == "网络安全"]
    grouped = {}
    for p in items:
        grouped.setdefault(p["category"], []).append(p["concept"])
    return "；".join(f"{cat}：{'、'.join(dict.fromkeys(cons))}" for cat, cons in grouped.items()) or "—"


wb = Workbook()
wb.remove(wb.active)
HEADER_FILL = PatternFill("solid", fgColor="122B49")
ALT_FILL = PatternFill("solid", fgColor="F5F5F5")
THIN = Side(style="thin", color="D0D0D0")
BLUE, GREEN, RED = "0066CC", "008000", "FF0000"
HEADERS = ["代码", "名称", "旧网络安全概念", "动作", "新一级分类", "新概念", "判词", "详细理由", "证据名称", "证据日期", "Source Name", "Source URL"]
WIDTHS = [12, 14, 38, 10, 16, 48, 10, 110, 42, 14, 52, 72]


def setup(ws, title, subtitle, ncols):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.row_dimensions[2].height = 32
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=1 + ncols)
    c = ws.cell(2, 2, title)
    c.font = Font(name="Microsoft YaHei", size=18, bold=True, color="122B49")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=1 + ncols)
    c = ws.cell(3, 2, subtitle)
    c.font = Font(name="Microsoft YaHei", size=10, color="666666")


def save_check(ws, min_rows):
    wb.save(OUT)
    check = load_workbook(OUT, data_only=False)
    target = check[ws.title]
    assert target.max_row >= min_rows and not target.sheet_view.showGridLines
    for row in target.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not any(e in cell.value for e in ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?"))
    check.close()


def add_sheet(name, title, subtitle, rows):
    ws = wb.create_sheet(name)
    setup(ws, title, subtitle, len(HEADERS))
    for j, h in enumerate(HEADERS, 2):
        c = ws.cell(4, j, h)
        c.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=THIN)
    for i, row in enumerate(rows, 5):
        pairs = row["pairs"]
        values = [row["code"], row["name"], pair_text(current[row["code"]]["pairs"], True), row["action"],
                  "、".join(p["category"] for p in pairs) or "—",
                  "；".join("、".join(p["concepts"]) for p in pairs) or "—",
                  row["verdict"], row["reason"], row["evidence_title"], row["evidence_date"],
                  f"iFinD ifind_get_stock_info (as of 2026-10-27)；{row['evidence_title']}", row["evidence_url"]]
        for j, value in enumerate(values, 2):
            c = ws.cell(i, j, value)
            c.font = Font(name="Microsoft YaHei", size=10, color=RED if j == 13 else BLUE)
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=j in (4, 7, 9, 12, 13))
            if i % 2 == 0:
                c.fill = ALT_FILL
        ws.row_dimensions[i].height = 92
    for j, width in enumerate(WIDTHS, 2):
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"B4:M{4 + len(rows)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    save_check(ws, 4 + len(rows))


kept = [x for x in audit if x["verdict"] != "摘除"]
removed = [x for x in audit if x["action"] == "摘除"]
add_sheet("清洗总表", "网络安全 · 45只逐项复核", "逐股核对主营、产品、客户与商业化证据；重写详细理由并清理泛化、错标标签。", audit)
add_sheet("网络安全", "网络安全 · 清洗后38只股池", "核心30只、相关8只；按实际产品与交付细化为47个二级概念。", kept)
add_sheet("摘除清单", "摘除清单 · 7只错标与落地不足", "普通信息化、内部平台治理、纸币鉴伪、传媒VR及参股关系不直接计入网络安全。", removed)

flat = wb.create_sheet("概念展开")
flat.append(["一级分类", "二级概念", "判词", "代码", "名称"])
for row in audit:
    for pair in row["pairs"]:
        for concept in pair["concepts"]:
            flat.append([pair["category"], concept, row["verdict"], row["code"], row["name"]])
flat.sheet_state = "hidden"
last_flat = flat.max_row

ws = wb.create_sheet("概念统计")
setup(ws, "概念统计 · 核心/相关分布", "公式引用隐藏展开表精确计数；一只股票可对应多个概念。", 5)
for j, h in enumerate(["一级分类", "二级概念", "核心（只）", "相关（只）", "合计（只）"], 2):
    c = ws.cell(4, j, h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center")
concepts = sorted({(p["category"], c) for x in audit for p in x["pairs"] for c in p["concepts"]})
for i, (category, concept) in enumerate(concepts, 5):
    ws.cell(i, 2, category)
    ws.cell(i, 3, concept).font = Font(color=BLUE)
    ws.cell(i, 4, f'=COUNTIFS(\'概念展开\'!$A$2:$A${last_flat},B{i},\'概念展开\'!$B$2:$B${last_flat},C{i},\'概念展开\'!$C$2:$C${last_flat},"核心")').font = Font(color=GREEN)
    ws.cell(i, 5, f'=COUNTIFS(\'概念展开\'!$A$2:$A${last_flat},B{i},\'概念展开\'!$B$2:$B${last_flat},C{i},\'概念展开\'!$C$2:$C${last_flat},"相关")').font = Font(color=GREEN)
    ws.cell(i, 6, f"=SUM(D{i}:E{i})")
    for j in range(2, 7):
        if i % 2 == 0:
            ws.cell(i, j).fill = ALT_FILL
last_stats = 4 + len(concepts)
ws.cell(last_stats + 1, 2, "合计（概念次数，非股票去重数）").font = Font(bold=True)
for j in range(4, 7):
    col = get_column_letter(j)
    ws.cell(last_stats + 1, j, f"=SUM({col}5:{col}{last_stats})").font = Font(bold=True)
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 34
for col in ["D", "E", "F"]:
    ws.column_dimensions[col].width = 18
ws.freeze_panes = "B5"
ws.auto_filter.ref = f"B4:F{last_stats}"
save_check(ws, last_stats + 1)

ws = wb.create_sheet("封面", 0)
setup(ws, "网络安全概念深度清洗报告", "复核日期：2026-10-27｜范围：45只｜重点：真实产品、商业化、分类边界与详细理由", 7)
for col, width in zip(range(2, 9), [30, 18, 18, 18, 18, 18, 18]):
    ws.column_dimensions[get_column_letter(col)].width = width
metrics = [
    ("复核股票（只）", f"=COUNTA('清洗总表'!$B$5:$B${4 + len(audit)})"),
    ("保留网络安全（只）", f"=COUNTA('网络安全'!$B$5:$B${4 + len(kept)})"),
    ("核心判词（只）", f'=COUNTIF(\'清洗总表\'!$H$5:$H${4 + len(audit)},"核心")'),
    ("相关判词（只）", f'=COUNTIF(\'清洗总表\'!$H$5:$H${4 + len(audit)},"相关")'),
    ("完全摘除（只）", f"=COUNTA('摘除清单'!$B$5:$B${4 + len(removed)})"),
    ("纠错动作（只）", f'=COUNTIF(\'清洗总表\'!$E$5:$E${4 + len(audit)},"纠错")'),
    ("新增动作（只）", f'=COUNTIF(\'清洗总表\'!$E$5:$E${4 + len(audit)},"新增")'),
]
for i, (label, formula) in enumerate(metrics, 6):
    ws.cell(i, 2, label).font = Font(name="Microsoft YaHei", bold=True, color="333333")
    ws.cell(i, 3, formula).font = Font(name="Microsoft YaHei", size=14, bold=True, color=GREEN)
    ws.row_dimensions[i].height = 26
ws.cell(15, 2, "工作表索引").font = Font(size=13, bold=True, color="122B49")
index_rows = [
    ("清洗总表", "45只旧概念、新概念、动作、判词、详细理由和证据"),
    ("网络安全", f"{len(kept)}只清洗后网络安全公司"),
    ("摘除清单", f"{len(removed)}只错标、内部治理或落地不足公司"),
    ("概念统计", f"{len(concepts)}个二级概念的核心/相关数量"),
]
for i, (name, desc) in enumerate(index_rows, 16):
    ws.cell(i, 2, name).font = Font(bold=True, color=BLUE)
    ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=7)
    ws.cell(i, 3, desc).font = Font(color=BLUE)
notes = [
    "核心：网络安全产品或服务是明确主营、重要收入或规模化核心产品；相关：有真实产品、客户、项目或收入，但仍属新业务或占比较低。",
    "普通安防、IDC托管、信息化集成、传媒VR、内部审核和纸币鉴伪不直接纳入；央企或国资背景也不再作为独立二级概念。",
    "“其他”“国家队”“字节生态”等泛标签已按防火墙、密码、数据安全、内容安全、安全运营、DPI等真实业务重新归类。",
    "主营资料：iFinD ifind_get_stock_info（2026-10-27）；逐股证据名称、日期及URL见Source列。",
]
for i, text in enumerate(notes, 23):
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=8)
    ws.cell(i, 2, text).font = Font(size=10, color="666666")
save_check(ws, 26)

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUT)
print(OUT)
print("sheets", wb.sheetnames)
print("rows", len(audit), len(kept), len(removed), len(concepts), last_flat - 1)
