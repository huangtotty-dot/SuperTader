# -*- coding: utf-8 -*-
"""军工清洗独立QA：交付数字直接由JSON重算，不依赖Excel公式缓存。"""
import json
import sys
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(r"E:\superTrader")
DC = ROOT / "data_cleaning"
SRC = ROOT / "stock_hunter" / "watchlist_jiuyan.json"
BAK = ROOT / "stock_hunter" / "watchlist_jiuyan.json.bak_20261027_military_deep"
PREVIEW = DC / "watchlist_jiuyan_military_cleaned_20261027.json"
XLSX = DC / "军工概念深度清洗_20261027.xlsx"
AUDIT = DC / "military_audit_20261027.json"
CURRENT = DC / "military_current_20261027.json"
OUT = DC / "military_validation_20261027.json"
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "stock_hunter"))
from modules.data_loader import WatchlistLoader

formal = json.loads(SRC.read_text(encoding="utf-8"))
backup = json.loads(BAK.read_text(encoding="utf-8"))
preview = json.loads(PREVIEW.read_text(encoding="utf-8"))
audit = json.loads(AUDIT.read_text(encoding="utf-8"))
current = json.loads(CURRENT.read_text(encoding="utf-8"))
target = {x["code"] for x in audit}
assert formal == preview and len(formal) == len(backup) == 5400
assert all(formal[c] == backup[c] for c in formal if c not in target)
assert len(audit) == len(target) == len(current) == 97
assert target == {x["code"] for x in current}
assert all(len(x["reason"]) >= 150 for x in audit)
assert all(x["evidence_title"] and x["evidence_date"] and x["evidence_url"] for x in audit)


def toks(value):
    if isinstance(value, list):
        return [str(x) for x in value if str(x)]
    return [x for x in str(value or "").split("|") if x]


def active_pairs(rec):
    out = []
    for i in list(range(1, 10)) + [0]:
        suffix = "" if i == 0 else str(i)
        for cat in toks(rec.get("jiuyan_category" + suffix)):
            for con in toks(rec.get("jiuyan_concept" + suffix)):
                if (cat, con) not in out:
                    out.append((cat, con))
    return out


military_rows = [x for x in audit if any(p["category"] == "军工" for p in x["pairs"])]
verdicts = Counter(x["verdict"] for x in military_rows)
actions = Counter(x["action"] for x in audit)
concepts = sorted({c for x in military_rows for p in x["pairs"] if p["category"] == "军工" for c in p["concepts"]})
concept_verdicts = Counter((p["category"], c, x["verdict"]) for x in audit for p in x["pairs"] for c in p["concepts"])
core_occurrences = sum(v for (_, _, verdict), v in concept_verdicts.items() if verdict == "核心")
related_occurrences = sum(v for (_, _, verdict), v in concept_verdicts.items() if verdict == "相关")
metrics = {
    "复核股票": len(audit),
    "保留军工": len(military_rows),
    "核心": verdicts["核心"],
    "相关": verdicts["相关"],
    "纠错": actions["纠错"],
    "完全摘除": actions["摘除"],
    "军工二级概念": len(concepts),
    "概念出现次数": core_occurrences + related_occurrences,
    "最短理由字符数": min(len(x["reason"]) for x in audit),
}
assert list(metrics.values())[:8] == [97, 89, 74, 15, 89, 8, 31, 150], metrics
old_concepts = {p["concept"] for x in current for p in x["pairs"] if p["category"] == "军工"}
removed_expected = {"002222", "002342", "002383", "002724", "002747", "300627", "600343", "600698"}
assert {x["code"] for x in audit if x["verdict"] == "摘除"} == removed_expected

for row in audit:
    code = row["code"]
    rec = formal[code]
    pairs = active_pairs(rec)
    expected = [(p["category"], c) for p in row["pairs"] for c in p["concepts"]]
    got = [pair for pair in pairs if pair[0] == "军工"]
    assert got == expected, (code, got, expected)
    assert not old_concepts.intersection(c for _, c in got)
    for _, concept in expected:
        assert rec["jiuyan_concept_reason"].get(concept) == row["reason"]
        assert rec["jiuyan_concept_desc"].get(concept)
    before_other = [pair for pair in active_pairs(backup[code]) if pair[0] != "军工"]
    after_other = [pair for pair in pairs if pair[0] != "军工"]
    assert before_other == after_other, (code, before_other, after_other)
    for key in backup[code]:
        if key.endswith("_verdicts") and key != "jiuyan_military_verdicts":
            assert rec.get(key) == backup[code].get(key), (code, key)

for code in ["002111", "300589", "688070"]:
    row = next(x for x in audit if x["code"] == code)
    assert row["verdict"] == "相关" and row.get("borderline_review") == "2026-10-27二次复核修正"
assert next(x for x in audit if x["code"] == "601606")["evidence_url"].endswith("id=12186260")

loader = WatchlistLoader()
after_df = loader.load(str(SRC))
before_df = loader.load(str(BAK))
after_counts = Counter(after_df["韭研分类"])
before_counts = Counter(before_df["韭研分类"])
category_differences = {c: [before_counts[c], after_counts[c]] for c in set(before_counts) | set(after_counts) if before_counts[c] != after_counts[c]}
assert category_differences["军工"] == [97, 89]
assert set(category_differences) == {"军工", ""}, category_differences
for category in ["云计算", "网络安全", "AI应用", "物理AI", "数据要素"]:
    assert before_counts[category] == after_counts[category]
assert after_counts["云计算"] == 37 and after_counts["网络安全"] == 38
assert after_counts["AI应用"] == 84 and after_counts["物理AI"] == 54 and after_counts["数据要素"] == 77
assert len(after_df.loc[after_df["韭研分类"] == "军工", "代码"]) == 89
loader_counts = {c: after_counts[c] for c in ["军工", "云计算", "网络安全", "AI应用", "物理AI", "数据要素"]}

wb = load_workbook(XLSX, data_only=False)
assert wb.sheetnames == ["封面", "清洗总表", "军工", "摘除清单", "概念展开", "概念统计"]
assert wb["概念展开"].sheet_state == "hidden"
assert wb["清洗总表"].max_row == 101 and wb["军工"].max_row == 93 and wb["摘除清单"].max_row == 12
assert wb["概念展开"].max_row == 151 and wb["概念统计"].max_row == 36
flat_rows = [tuple(cell.value for cell in row) for row in wb["概念展开"].iter_rows(min_row=2, max_col=5)]
expected_flat = [(p["category"], c, x["verdict"], x["code"], x["name"]) for x in audit for p in x["pairs"] for c in p["concepts"]]
assert flat_rows == expected_flat
stats = {}
for row in wb["概念统计"].iter_rows(min_row=5, max_row=35, min_col=2, max_col=3, values_only=True):
    stats[row] = (concept_verdicts[(row[0], row[1], "核心")], concept_verdicts[(row[0], row[1], "相关")])
assert len(stats) == 31 and {con for _, con in stats} == set(concepts)
assert sum(v[0] for v in stats.values()) == core_occurrences
assert sum(v[1] for v in stats.values()) == related_occurrences
formula_count = sum(1 for ws in wb for row in ws.iter_rows() for cell in row if cell.data_type == "f")
assert formula_count == 102
assert wb.calculation.calcMode == "auto" and wb.calculation.fullCalcOnLoad and wb.calculation.forceFullCalc
wb.close()

result = {
    "status": "pass_independent_qa",
    "metrics": metrics,
    "concept_occurrences_by_verdict": {"核心": core_occurrences, "相关": related_occurrences, "合计": core_occurrences + related_occurrences},
    "formal_equals_preview": True,
    "backup_records": len(backup),
    "non_target_records_unchanged": len(formal) - len(target),
    "other_category_pairs_and_verdicts_preserved": True,
    "category_differences_loader": category_differences,
    "loader_counts": loader_counts,
    "borderline_second_review": {"002111": "相关", "300589": "相关", "688070": "相关"},
    "workbook_formula_count": formula_count,
    "workbook_formula_gate": "deferred_calcMode_auto; 102 formula caches unavailable; all reported values independently recomputed from JSON",
    "workbook_independent_stats": {f"{cat}/{con}": {"核心": counts[0], "相关": counts[1]} for (cat, con), counts in stats.items()},
}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
