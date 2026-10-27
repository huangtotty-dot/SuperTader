# -*- coding: utf-8 -*-
"""商业航天清洗独立QA：交付数字直接由JSON重算，不依赖Excel公式缓存。"""
import json
import sys
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(r"E:\superTrader")
DC = ROOT / "data_cleaning"
SRC = ROOT / "stock_hunter" / "watchlist_jiuyan.json"
BAK = ROOT / "stock_hunter" / "watchlist_jiuyan.json.bak_20261027_commercial_space_deep"
PREVIEW = DC / "watchlist_jiuyan_commercial_space_cleaned_20261027.json"
XLSX = DC / "商业航天概念深度清洗_20261027.xlsx"
AUDIT = DC / "commercial_space_audit_20261027.json"
CURRENT = DC / "commercial_space_current_20261027.json"
REVIEW = DC / "commercial_space_borderline_review_20261027.json"
OUT = DC / "commercial_space_validation_20261027.json"
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "stock_hunter"))
from modules.data_loader import WatchlistLoader

formal = json.loads(SRC.read_text(encoding="utf-8"))
backup = json.loads(BAK.read_text(encoding="utf-8"))
preview = json.loads(PREVIEW.read_text(encoding="utf-8"))
audit = json.loads(AUDIT.read_text(encoding="utf-8"))
current = json.loads(CURRENT.read_text(encoding="utf-8"))
review = json.loads(REVIEW.read_text(encoding="utf-8"))
target = {x["code"] for x in audit}
assert formal == preview and len(formal) == len(backup) == 5400
assert all(formal[c] == backup[c] for c in formal if c not in target)
assert len(audit) == len(target) == len(current) == 90
assert target == {x["code"] for x in current}
assert all(len(x["reason"]) >= 150 for x in audit)
assert all(x["evidence_title"] and x["evidence_date"] and x["evidence_url"] for x in audit)
assert all(x["pairs"] == [] for x in audit if x["verdict"] == "摘除")
assert all(x["pairs"] and all(p["category"] == "商业航天" for p in x["pairs"]) for x in audit if x["verdict"] != "摘除")


def toks(value):
    if isinstance(value, list):
        return [str(x) for x in value if str(x)]
    return [x for x in str(value or "").split("|") if x]


def active_pairs(rec):
    out = []
    for i in list(range(1, 10)) + [0]:
        suffix = "" if i == 0 else str(i)
        for cat in toks(rec.get("jiuyan_category" + suffix)):
            for con in toks(rec.get("jiuyan_concept" + suffix)):
                if (cat, con) not in out:
                    out.append((cat, con))
    return out


commercial_rows = [x for x in audit if any(p["category"] == "商业航天" for p in x["pairs"])]
verdicts = Counter(x["verdict"] for x in commercial_rows)
actions = Counter(x["action"] for x in audit)
concepts = sorted({c for x in commercial_rows for p in x["pairs"] if p["category"] == "商业航天" for c in p["concepts"]})
concept_verdicts = Counter((p["category"], c, x["verdict"]) for x in audit for p in x["pairs"] for c in p["concepts"])
core_occurrences = sum(v for (_, _, verdict), v in concept_verdicts.items() if verdict == "核心")
related_occurrences = sum(v for (_, _, verdict), v in concept_verdicts.items() if verdict == "相关")
metrics = {
    "复核股票": len(audit),
    "保留商业航天": len(commercial_rows),
    "核心": verdicts["核心"],
    "相关": verdicts["相关"],
    "纠错": actions["纠错"],
    "完全摘除": actions["摘除"],
    "商业航天二级概念": len(concepts),
    "概念出现次数": core_occurrences + related_occurrences,
    "最短理由字符数": min(len(x["reason"]) for x in audit),
}
assert list(metrics.values())[:8] == [90, 74, 14, 60, 74, 16, 23, 111], metrics
old_concepts = {p["concept"] for x in current for p in x["pairs"] if p["category"] == "商业航天"}
removed_expected = {"001379", "002088", "002131", "002245", "002796", "002985", "300118", "300316", "300604", "300751", "300757", "301071", "600456", "600855", "688516", "688610"}
assert {x["code"] for x in audit if x["verdict"] == "摘除"} == removed_expected

for row in audit:
    code = row["code"]
    rec = formal[code]
    pairs = active_pairs(rec)
    expected = [(p["category"], c) for p in row["pairs"] for c in p["concepts"]]
    got = [pair for pair in pairs if pair[0] == "商业航天"]
    assert got == expected, (code, got, expected)
    assert not old_concepts.intersection(c for _, c in got)
    for _, concept in expected:
        assert rec["jiuyan_concept_reason"].get(concept) == row["reason"]
        assert rec["jiuyan_concept_desc"].get(concept)
    before_other = [pair for pair in active_pairs(backup[code]) if pair[0] != "商业航天"]
    after_other = [pair for pair in pairs if pair[0] != "商业航天"]
    assert before_other == after_other, (code, before_other, after_other)
    for key in backup[code]:
        if key.endswith("_verdicts") and key != "jiuyan_commercial_space_verdicts":
            assert rec.get(key) == backup[code].get(key), (code, key)
    audit_mark = rec["jiuyan_commercial_space_audit_20261027"]
    assert rec["jiuyan_commercial_space_audited_at"] == "2026-10-27"
    assert audit_mark["reason"] == row["reason"] and audit_mark["verdict"] == row["verdict"]
    assert rec["jiuyan_commercial_space_evidence"]["url"] == row["evidence_url"]
    if row["verdict"] == "摘除":
        assert "jiuyan_commercial_space_verdicts" not in rec
    else:
        assert rec["jiuyan_commercial_space_verdicts"] == {"商业航天": row["verdict"]}

for code, verdict in [("600879", "相关"), ("603809", "相关")]:
    row = next(x for x in audit if x["code"] == code)
    assert row["verdict"] == verdict and row.get("borderline_review") == "2026-10-27二次复核修正"
assert next(x for x in audit if x["code"] == "600879")["evidence_url"].endswith("AN202512151801065533.html")
assert next(x for x in audit if x["code"] == "603809")["evidence_url"].endswith("AN202601121816949364.html")
assert review["summary"]["verdict_change_count"] == 2

loader = WatchlistLoader()
after_df = loader.load(str(SRC))
before_df = loader.load(str(BAK))
after_counts = Counter(after_df["韭研分类"])
before_counts = Counter(before_df["韭研分类"])
category_differences = {c: [before_counts[c], after_counts[c]] for c in set(before_counts) | set(after_counts) if before_counts[c] != after_counts[c]}
assert category_differences["商业航天"] == [90, 74]
assert set(category_differences) == {"商业航天", ""}, category_differences
for category in ["军工", "云计算", "网络安全", "AI应用", "物理AI", "数据要素"]:
    assert before_counts[category] == after_counts[category]
assert after_counts["军工"] == 89 and after_counts["云计算"] == 37 and after_counts["网络安全"] == 38
assert after_counts["AI应用"] == 84 and after_counts["物理AI"] == 54 and after_counts["数据要素"] == 77
assert len(after_df.loc[after_df["韭研分类"] == "商业航天", "代码"]) == 74
loader_counts = {c: after_counts[c] for c in ["商业航天", "军工", "云计算", "网络安全", "AI应用", "物理AI", "数据要素"]}

wb = load_workbook(XLSX, data_only=False)
assert wb.sheetnames == ["封面", "清洗总表", "商业航天", "摘除清单", "概念展开", "概念统计"]
assert wb["概念展开"].sheet_state == "hidden"
assert wb["清洗总表"].max_row == 94 and wb["商业航天"].max_row == 78 and wb["摘除清单"].max_row == 20
assert wb["概念展开"].max_row == 112 and wb["概念统计"].max_row == 28
flat_rows = [tuple(cell.value for cell in row) for row in wb["概念展开"].iter_rows(min_row=2, max_col=5)]
expected_flat = [(p["category"], c, x["verdict"], x["code"], x["name"]) for x in audit for p in x["pairs"] for c in p["concepts"]]
assert flat_rows == expected_flat
stats = {}
for row in wb["概念统计"].iter_rows(min_row=5, max_row=27, min_col=2, max_col=3, values_only=True):
    stats[row] = (concept_verdicts[(row[0], row[1], "核心")], concept_verdicts[(row[0], row[1], "相关")])
assert len(stats) == 23 and {con for _, con in stats} == set(concepts)
assert sum(v[0] for v in stats.values()) == core_occurrences
assert sum(v[1] for v in stats.values()) == related_occurrences
formula_count = sum(1 for ws in wb for row in ws.iter_rows() for cell in row if cell.data_type == "f")
assert formula_count == 78
assert wb.calculation.calcMode == "auto" and wb.calculation.fullCalcOnLoad and wb.calculation.forceFullCalc
wb.close()

result = {
    "status": "pass_independent_qa",
    "metrics": metrics,
    "concept_occurrences_by_verdict": {"核心": core_occurrences, "相关": related_occurrences, "合计": core_occurrences + related_occurrences},
    "formal_equals_preview": True,
    "backup_records": len(backup),
    "non_target_records_unchanged": len(formal) - len(target),
    "other_category_pairs_and_verdicts_preserved": True,
    "category_differences_loader": category_differences,
    "loader_counts": loader_counts,
    "borderline_second_review": {"600879": "核心→相关", "603809": "摘除→相关"},
    "workbook_formula_count": formula_count,
    "workbook_formula_gate": "deferred_calcMode_auto; 78 formula caches unavailable; all reported values independently recomputed from JSON",
    "workbook_independent_stats": {f"{cat}/{con}": {"核心": counts[0], "相关": counts[1]} for (cat, con), counts in stats.items()},
}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
