# -*- coding: utf-8 -*-
from pathlib import Path
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"E:\superTrader\data_cleaning")
OUT = ROOT / "军工概念深度清洗_20261027.xlsx"
audit = json.loads((ROOT / "military_audit_20261027.json").read_text(encoding="utf-8"))
current = {x["code"]: x for x in json.loads((ROOT / "military_current_20261027.json").read_text(encoding="utf-8"))}


def pair_text(pairs, military_only=False):
    grouped = {}
    for p in pairs:
        if military_only and p["category"] != "军工":
            continue
        grouped.setdefault(p["category"], []).append(p["concept"])
    return "；".join(f"{cat}：{'、'.join(dict.fromkeys(cons))}" for cat, cons in grouped.items()) or "—"


wb = Workbook()
wb.remove(wb.active)
HEADER_FILL = PatternFill("solid", fgColor="263A2D")
ALT_FILL = PatternFill("solid", fgColor="F5F5F5")
THIN = Side(style="thin", color="D0D0D0")
BLUE, GREEN, RED = "0066CC", "008000", "FF0000"
HEADERS = ["代码", "名称", "旧军工概念", "动作", "新一级分类", "新概念", "判词", "详细理由", "证据名称", "证据日期", "Source Name", "Source URL"]
WIDTHS = [12, 14, 36, 10, 16, 55, 10, 115, 44, 14, 62, 74]


def setup(ws, title, subtitle, ncols):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.row_dimensions[2].height = 32
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=1 + ncols)
    c = ws.cell(2, 2, title)
    c.font = Font(name="Microsoft YaHei", size=18, bold=True, color="263A2D")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=1 + ncols)
    c = ws.cell(3, 2, subtitle)
    c.font = Font(name="Microsoft YaHei", size=10, color="666666")


def save_check(ws, min_rows):
    wb.save(OUT)
    check = load_workbook(OUT, data_only=False)
    target = check[ws.title]
    assert target.max_row >= min_rows and not target.sheet_view.showGridLines
    for row in target.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not any(e in cell.value for e in ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?"))
    check.close()


def add_sheet(name, title, subtitle, rows, mode="all"):
    ws = wb.create_sheet(name)
    setup(ws, title, subtitle, len(HEADERS))
    for j, h in enumerate(HEADERS, 2):
        c = ws.cell(4, j, h)
        c.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=THIN)
    for i, row in enumerate(rows, 5):
        pairs = row["pairs"]
        if mode == "military":
            pairs = [p for p in pairs if p["category"] == "军工"]
        values = [row["code"], row["name"], pair_text(current[row["code"]]["pairs"], True), row["action"],
                  "、".join(p["category"] for p in pairs) or "—",
                  "；".join("、".join(p["concepts"]) for p in pairs) or "—",
                  row["verdict"], row["reason"], row["evidence_title"], row["evidence_date"],
                  f"iFinD ifind_get_stock_info + ifind_get_stock_business_segmentation (as of 2026-10-27)；{row['evidence_title']}", row["evidence_url"]]
        for j, value in enumerate(values, 2):
            c = ws.cell(i, j, value)
            c.font = Font(name="Microsoft YaHei", size=10, color=RED if j == 13 else BLUE)
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=j in (4, 7, 9, 12, 13))
            if i % 2 == 0:
                c.fill = ALT_FILL
        ws.row_dimensions[i].height = 96
    for j, width in enumerate(WIDTHS, 2):
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"B4:M{4 + len(rows)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    save_check(ws, 4 + len(rows))


military = [x for x in audit if any(p["category"] == "军工" for p in x["pairs"])]
removed = [x for x in audit if x["action"] == "摘除"]
add_sheet("清洗总表", "军工 · 97只逐项复核", "逐股核对军工产品、军用场景、客户与收入；区分真实军工配套和通用产品联想。", audit)
add_sheet("军工", "军工 · 清洗后89只股池", "核心74只、相关15只；按装备、分系统、材料和软件细化为31个二级概念。", military, "military")
add_sheet("摘除清单", "摘除清单 · 8只证据不足或业务错标", "通用光电、索具、民用导航、工业机器人及集团背景不直接计入军工。", removed)

flat = wb.create_sheet("概念展开")
flat.append(["一级分类", "二级概念", "判词", "代码", "名称"])
for row in audit:
    for pair in row["pairs"]:
        for concept in pair["concepts"]:
            flat.append([pair["category"], concept, row["verdict"], row["code"], row["name"]])
flat.sheet_state = "hidden"
last_flat = flat.max_row

ws = wb.create_sheet("概念统计")
setup(ws, "概念统计 · 核心/相关分布", "公式引用隐藏展开表精确计数；一只股票可对应多个概念。", 5)
for j, h in enumerate(["一级分类", "二级概念", "核心（只）", "相关（只）", "合计（只）"], 2):
    c = ws.cell(4, j, h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center")
concepts = sorted({(p["category"], c) for x in audit for p in x["pairs"] for c in p["concepts"]})
for i, (category, concept) in enumerate(concepts, 5):
    ws.cell(i, 2, category)
    ws.cell(i, 3, concept).font = Font(color=BLUE)
    ws.cell(i, 4, f'=COUNTIFS(\'概念展开\'!$A$2:$A${last_flat},B{i},\'概念展开\'!$B$2:$B${last_flat},C{i},\'概念展开\'!$C$2:$C${last_flat},"核心")').font = Font(color=GREEN)
    ws.cell(i, 5, f'=COUNTIFS(\'概念展开\'!$A$2:$A${last_flat},B{i},\'概念展开\'!$B$2:$B${last_flat},C{i},\'概念展开\'!$C$2:$C${last_flat},"相关")').font = Font(color=GREEN)
    ws.cell(i, 6, f"=SUM(D{i}:E{i})")
    for j in range(2, 7):
        if i % 2 == 0:
            ws.cell(i, j).fill = ALT_FILL
last_stats = 4 + len(concepts)
ws.cell(last_stats + 1, 2, "合计（概念次数，非股票去重数）").font = Font(bold=True)
for j in range(4, 7):
    col = get_column_letter(j)
    ws.cell(last_stats + 1, j, f"=SUM({col}5:{col}{last_stats})").font = Font(bold=True)
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 42
for col in ["D", "E", "F"]:
    ws.column_dimensions[col].width = 18
ws.freeze_panes = "B5"
ws.auto_filter.ref = f"B4:F{last_stats}"
save_check(ws, last_stats + 1)

ws = wb.create_sheet("封面", 0)
setup(ws, "军工概念深度清洗报告", "复核日期：2026-10-27｜范围：97只｜重点：军工实质、产品层级、军民边界与二次复审", 7)
for col, width in zip(range(2, 9), [30, 18, 18, 18, 18, 18, 18]):
    ws.column_dimensions[get_column_letter(col)].width = width
metrics = [
    ("复核股票（只）", f"=COUNTA('清洗总表'!$B$5:$B${4 + len(audit)})"),
    ("保留军工（只）", f"=COUNTA('军工'!$B$5:$B${4 + len(military)})"),
    ("核心判词（只）", f'=COUNTIF(\'军工\'!$H$5:$H${4 + len(military)},"核心")'),
    ("相关判词（只）", f'=COUNTIF(\'军工\'!$H$5:$H${4 + len(military)},"相关")'),
    ("纠错动作（只）", f'=COUNTIF(\'清洗总表\'!$E$5:$E${4 + len(audit)},"纠错")'),
    ("完全摘除（只）", f"=COUNTA('摘除清单'!$B$5:$B${4 + len(removed)})"),
]
for i, (label, formula) in enumerate(metrics, 6):
    ws.cell(i, 2, label).font = Font(name="Microsoft YaHei", bold=True, color="333333")
    ws.cell(i, 3, formula).font = Font(name="Microsoft YaHei", size=14, bold=True, color=GREEN)
    ws.row_dimensions[i].height = 26
ws.cell(15, 2, "工作表索引").font = Font(size=13, bold=True, color="263A2D")
index_rows = [
    ("清洗总表", "97只旧概念、新概念、动作、判词、详细理由和证据"),
    ("军工", f"{len(military)}只清洗后军工公司"),
    ("摘除清单", f"{len(removed)}只通用业务或军工证据不足公司"),
    ("概念统计", f"{len(concepts)}个二级概念"),
]
for i, (name, desc) in enumerate(index_rows, 16):
    ws.cell(i, 2, name).font = Font(bold=True, color=BLUE)
    ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=7)
    ws.cell(i, 3, desc).font = Font(color=BLUE)
notes = [
    "核心：军工是明确主营、重要收入来源，或已有规模化核心产品；相关：有真实军工产品、订单或交付，但民用为主或收入未独立拆分。",
    "摘除：只有集团背景、军工资质、潜在供应、通用产品可军用，或当前只有民用机器人、导航、材料等联想。",
    "二次复审：威海广泰、江龙船艇、纵横股份依据真实历史订单或2025年军工交付证据，由初审摘除修正为相关。",
    "主营与2025分部：iFinD ifind_get_stock_info、ifind_get_stock_business_segmentation（2026-10-27）；逐股证据见Source列。",
]
for i, text in enumerate(notes, 23):
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=8)
    ws.cell(i, 2, text).font = Font(size=10, color="666666")
save_check(ws, 26)

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUT)
print(OUT)
print("sheets", wb.sheetnames)
print("rows", len(audit), len(military), len(removed), len(concepts), last_flat - 1)
