# -*- coding: utf-8 -*-
from pathlib import Path
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"E:\superTrader\data_cleaning")
OUT = ROOT / "数据要素概念深度清洗_20261027.xlsx"
audit = json.loads((ROOT / "dataelement_audit_20261027.json").read_text(encoding="utf-8"))
current = {x["code"]: x for x in json.loads((ROOT / "dataelement_current.json").read_text(encoding="utf-8"))}
for row in audit:
    row["old_concepts"] = "、".join(dict.fromkeys(current[row["code"]]["concepts"]))

wb = Workbook()
wb.remove(wb.active)
HEADERS = ["代码", "名称", "旧概念", "动作", "新一级分类", "新概念", "判词", "详细理由", "证据名称", "证据日期", "Source Name", "Source URL"]
WIDTHS = [12, 14, 32, 16, 14, 28, 10, 95, 36, 14, 46, 72]
HEADER_FILL = PatternFill("solid", fgColor="122B49")
ALT_FILL = PatternFill("solid", fgColor="F5F5F5")
BLUE = "0066CC"
GREEN = "008000"
RED = "FF0000"
THIN = Side(style="thin", color="D0D0D0")


def keys(row):
    return [row["code"], row["name"], row["old_concepts"], row["action"], row["category"], "、".join(row["concepts"]), row["verdict"], row["reason"], row["evidence_title"], row["evidence_date"], f"iFinD ifind_get_stock_info (as of 2026-10-27)；{row['evidence_title']}", row["evidence_url"]]


def setup(ws, title, subtitle, ncols):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.row_dimensions[2].height = 32
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=1+ncols)
    c = ws.cell(2, 2, title)
    c.font = Font(name="Microsoft YaHei", size=18, bold=True, color="122B49")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=1+ncols)
    c = ws.cell(3, 2, subtitle)
    c.font = Font(name="Microsoft YaHei", size=10, color="666666")
    c.alignment = Alignment(horizontal="left", vertical="center")


def save_check(ws, min_rows):
    wb.save(OUT)
    check = load_workbook(OUT, read_only=False, data_only=False)
    assert ws.title in check.sheetnames
    target = check[ws.title]
    assert target.max_row >= min_rows, (ws.title, target.max_row, min_rows)
    assert target.sheet_view.showGridLines is False
    for row in target.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert "#REF!" not in cell.value and "#VALUE!" not in cell.value
    check.close()


def add_data_sheet(name, title, subtitle, rows):
    ws = wb.create_sheet(name)
    setup(ws, title, subtitle, len(HEADERS))
    for j, h in enumerate(HEADERS, 2):
        c = ws.cell(4, j, h)
        c.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        c.border = Border(bottom=THIN)
    ws.row_dimensions[4].height = 26
    for i, row in enumerate(rows, 5):
        values = keys(row)
        for j, value in enumerate(values, 2):
            c = ws.cell(i, j, value)
            c.font = Font(name="Microsoft YaHei", size=10, color=RED if j == 13 else BLUE)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
            if i % 2 == 0:
                c.fill = ALT_FILL
        ws.row_dimensions[i].height = 24
    for j, width in enumerate(WIDTHS, 2):
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"B4:M{4+len(rows)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    save_check(ws, 4+len(rows))
    return ws


# 1. 全量清洗表
add_data_sheet("清洗总表", "数据要素概念深度清洗 · 全量明细", "98 只逐项复核；理由与二级概念同步重写，证据可点击或复制访问。", audit)
# 2. 严格口径保留股池
kept = [x for x in audit if x["category"] == "数据要素"]
add_data_sheet("保留股池", "数据要素 · 严格口径保留股池", "仅保留核心与相关标的；二级概念已按实际产品、平台和交付方式重新归类。", kept)
# 3. 摘除与迁出
removed = [x for x in audit if x["category"] != "数据要素"]
add_data_sheet("摘除与迁出", "摘除与迁出 · 审计清单", "完全摘除 14 只；迁回物理 AI 7 只。历史标签仍保留在 watchlist 审计字段中。", removed)
# 4. 数据存储专项（截图中的原15只）
storage_codes = ["002368","601728","000948","603138","600050","600941","000032","000158","000815","000938","300895","002642","002152","600602","600037"]
storage = [next(x for x in audit if x["code"] == code) for code in storage_codes]
add_data_sheet("数据存储复核", "数据存储 · 截图原15只专项复核", "重点纠正把数据库、数据治理、数据运营和IDC全部混入“数据存储”的旧口径。", storage)

# 5. 概念统计（数值全部由Excel公式计算）
ws = wb.create_sheet("概念统计")
setup(ws, "保留股池 · 二级概念分布", "统计范围仅为清洗后保留的77只；数量由公式引用“保留股池”计算。", 4)
for j, h in enumerate(["二级概念", "核心（只）", "相关（只）", "合计（只）"], 2):
    c = ws.cell(4, j, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = HEADER_FILL; c.alignment = Alignment(horizontal="center")
concepts = sorted({c for x in kept for c in x["concepts"]})
last_kept = 4 + len(kept)
for i, concept in enumerate(concepts, 5):
    ws.cell(i, 2, concept).font = Font(color=BLUE)
    ws.cell(i, 3, f'=COUNTIFS(\'保留股池\'!$G$5:$G${last_kept},B{i},\'保留股池\'!$H$5:$H${last_kept},"核心")').font = Font(color=GREEN)
    ws.cell(i, 4, f'=COUNTIFS(\'保留股池\'!$G$5:$G${last_kept},B{i},\'保留股池\'!$H$5:$H${last_kept},"相关")').font = Font(color=GREEN)
    ws.cell(i, 5, f"=SUM(C{i}:D{i})").font = Font(color="000000")
    for j in range(2, 6):
        ws.cell(i, j).alignment = Alignment(horizontal="left" if j == 2 else "center", vertical="center", wrap_text=False)
        if i % 2 == 0: ws.cell(i, j).fill = ALT_FILL
last = 4 + len(concepts)
ws.cell(last+1, 2, "合计").font = Font(bold=True)
for j in range(3, 6):
    col = get_column_letter(j)
    ws.cell(last+1, j, f"=SUM({col}5:{col}{last})").font = Font(bold=True, color="000000")
ws.column_dimensions["B"].width = 34
for col in ["C","D","E"]: ws.column_dimensions[col].width = 16
ws.freeze_panes = "B5"
ws.auto_filter.ref = f"B4:E{last}"
save_check(ws, last+1)

# 6. 封面（最后创建，移动到首位）
ws = wb.create_sheet("封面", 0)
setup(ws, "数据要素概念深度清洗报告", "复核日期：2026-10-27｜范围：原98只候选｜口径：核心/相关保留，边缘/错标退出，物理AI单列", 7)
for col, width in zip(range(2, 9), [24,18,18,18,18,18,18]): ws.column_dimensions[get_column_letter(col)].width = width
metrics = [
    ("原候选股数（只）", "=COUNTA('清洗总表'!$B$5:$B$102)"),
    ("保留数据要素（只）", f"=COUNTA('保留股池'!$B$5:$B${last_kept})"),
    ("其中：核心（只）", f'=COUNTIF(\'保留股池\'!$H$5:$H${last_kept},"核心")'),
    ("其中：相关（只）", f'=COUNTIF(\'保留股池\'!$H$5:$H${last_kept},"相关")'),
    ("迁回物理AI（只）", "=COUNTIF('摘除与迁出'!$F$5:$F$25,\"物理AI\")"),
    ("完全摘除（只）", "=COUNTIF('摘除与迁出'!$E$5:$E$25,\"摘除\")"),
]
for i, (label, formula) in enumerate(metrics, 6):
    ws.cell(i, 2, label).font = Font(name="Microsoft YaHei", bold=True, color="333333")
    ws.cell(i, 3, formula).font = Font(name="Microsoft YaHei", size=14, bold=True, color=GREEN)
    ws.cell(i, 3).alignment = Alignment(horizontal="center")
    ws.row_dimensions[i].height = 26
ws.cell(14, 2, "工作表索引").font = Font(size=13, bold=True, color="122B49")
index_rows = [
    ("清洗总表", "98只全量旧概念、新概念、判词、详细理由及证据"),
    ("保留股池", "清洗后77只核心/相关标的"),
    ("摘除与迁出", "14只摘除及7只物理AI迁移的审计理由"),
    ("数据存储复核", "截图原15只的专项纠正"),
    ("概念统计", "保留股池二级概念的核心/相关数量"),
]
for i, (name, desc) in enumerate(index_rows, 15):
    ws.cell(i, 2, name).font = Font(bold=True, color=BLUE)
    ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=7)
    ws.cell(i, 3, desc).font = Font(color=BLUE)
    ws.cell(i, 3).alignment = Alignment(horizontal="left", wrap_text=False)
notes = [
    "判定口径：核心=该数据产品/服务是明确主要产品；相关=有真实产品或项目但主要仍为信息化、云、硬件或项目交付。",
    "退出口径：边缘=少量参股、自用数据或上游感知设备；错标=无直接证据、业务已失效或二级概念明显不符。",
    "主营资料：iFinD ifind_get_stock_info（2026-10-27）；概念证据：逐股年报、公告、公司官网或互动平台，详见Source列。",
]
for i, text in enumerate(notes, 22):
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=8)
    ws.cell(i, 2, text).font = Font(size=10, color="666666")
    ws.cell(i, 2).alignment = Alignment(horizontal="left", wrap_text=False)
save_check(ws, 24)

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUT)
print(OUT)
print("sheets", wb.sheetnames)
print("rows", len(audit), len(kept), len(removed), len(storage), len(concepts))
