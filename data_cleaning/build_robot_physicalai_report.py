# -*- coding: utf-8 -*-
from pathlib import Path
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"E:\superTrader\data_cleaning")
OUT = ROOT / "机器人与物理AI概念深度清洗_20261027.xlsx"
audit = json.loads((ROOT / "robot_physicalai_audit_20261027.json").read_text(encoding="utf-8"))
current = {x["code"]: x for x in json.loads((ROOT / "robot_current_20261027.json").read_text(encoding="utf-8"))}


def old_pairs(code):
    return "；".join(f"{p['category']}：{'、'.join(p['concepts'])}" for p in current[code]["pairs"]) or "—"


def new_pairs(row, category=None):
    pairs = [p for p in row["pairs"] if category is None or p["category"] == category]
    return "；".join(f"{p['category']}：{'、'.join(p['concepts'])}" for p in pairs) or "—"


def verdict(row, category=None):
    if category and row["code"] in {"300607", "688306"} and len(row["pairs"]) > 1:
        return "核心" if category == "机器人" else "相关"
    return row["verdict"]


wb = Workbook()
wb.remove(wb.active)
HEADER_FILL = PatternFill("solid", fgColor="122B49")
ALT_FILL = PatternFill("solid", fgColor="F5F5F5")
THIN = Side(style="thin", color="D0D0D0")
BLUE, GREEN, RED = "0066CC", "008000", "FF0000"
HEADERS = ["代码", "名称", "旧分类与概念", "动作", "新一级分类", "新概念", "判词", "详细理由", "证据名称", "证据日期", "Source Name", "Source URL"]
WIDTHS = [12, 14, 48, 12, 16, 42, 10, 105, 42, 14, 50, 72]


def setup(ws, title, subtitle, ncols):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.row_dimensions[2].height = 32
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=1 + ncols)
    c = ws.cell(2, 2, title)
    c.font = Font(name="Microsoft YaHei", size=18, bold=True, color="122B49")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=1 + ncols)
    c = ws.cell(3, 2, subtitle)
    c.font = Font(name="Microsoft YaHei", size=10, color="666666")
    c.alignment = Alignment(horizontal="left", vertical="center")


def save_check(ws, min_rows):
    wb.save(OUT)
    check = load_workbook(OUT, read_only=False, data_only=False)
    assert ws.title in check.sheetnames
    target = check[ws.title]
    assert target.max_row >= min_rows
    assert target.sheet_view.showGridLines is False
    for row in target.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not any(e in cell.value for e in ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?"))
    check.close()


def add_data_sheet(name, title, subtitle, rows, category=None, removed=False):
    ws = wb.create_sheet(name)
    setup(ws, title, subtitle, len(HEADERS))
    for j, h in enumerate(HEADERS, 2):
        c = ws.cell(4, j, h)
        c.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(bottom=THIN)
    ws.row_dimensions[4].height = 26
    for i, row in enumerate(rows, 5):
        pairs = [p for p in row["pairs"] if category is None or p["category"] == category]
        cats = "、".join(p["category"] for p in pairs) or "—"
        cons = "；".join("、".join(p["concepts"]) for p in pairs) or "—"
        values = [row["code"], row["name"], old_pairs(row["code"]), row["action"], cats, cons,
                  verdict(row, category), row["reason"], row["evidence_title"], row["evidence_date"],
                  f"iFinD ifind_get_stock_info (as of 2026-10-27)；{row['evidence_title']}", row["evidence_url"]]
        for j, value in enumerate(values, 2):
            c = ws.cell(i, j, value)
            c.font = Font(name="Microsoft YaHei", size=10, color=RED if j == 13 else BLUE)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
            if i % 2 == 0:
                c.fill = ALT_FILL
        ws.row_dimensions[i].height = 24
    for j, width in enumerate(WIDTHS, 2):
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"B4:M{4 + len(rows)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    save_check(ws, 4 + len(rows))
    return ws


add_data_sheet("清洗总表", "机器人与物理AI · 137只逐项复核", "每只股票一行；保留双分类，详细说明产品关系、商业化阶段及核心/相关判词。", audit)
robot = [x for x in audit if any(p["category"] == "机器人" for p in x["pairs"])]
physical = [x for x in audit if any(p["category"] == "物理AI" for p in x["pairs"])]
removed = [x for x in audit if x["action"] == "摘除"]
add_data_sheet("机器人", "机器人 · 保留与纠错股池", "硬件零部件不因人形机器人应用自动迁入物理AI；系统集成、本体与零部件分开。", robot, "机器人")
add_data_sheet("物理AI", "物理AI · 感知、控制、训练与具身本体", "统一具身本体/传感器口径；传统工业机器人收入不直接作为具身智能核心依据。", physical, "物理AI")
add_data_sheet("摘除清单", "摘除清单 · 错标、间接映射与研发预期", "仅参股、通用能力映射、无直接产品证据或仍处研发送样且无商业化的标签退出。", removed, removed=True)

# 隐藏展开表：统计公式按精确概念计数，避免“电机”误中“微特减速电机”。
flat = wb.create_sheet("概念展开")
flat.append(["一级分类", "二级概念", "判词", "代码", "名称"])
for row in audit:
    for pair in row["pairs"]:
        for concept in pair["concepts"]:
            flat.append([pair["category"], concept, verdict(row, pair["category"]), row["code"], row["name"]])
flat.sheet_state = "hidden"
last_flat = flat.max_row

ws = wb.create_sheet("概念统计")
setup(ws, "概念统计 · 核心/相关分布", "公式引用隐藏的概念展开表精确计数；一只股票可对应多个概念。", 5)
for j, h in enumerate(["一级分类", "二级概念", "核心（只）", "相关（只）", "合计（只）"], 2):
    c = ws.cell(4, j, h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center")
concepts = sorted({(p["category"], c) for x in audit for p in x["pairs"] for c in p["concepts"]})
for i, (category, concept) in enumerate(concepts, 5):
    ws.cell(i, 2, category)
    ws.cell(i, 3, concept).font = Font(color=BLUE)
    ws.cell(i, 4, f'=COUNTIFS(\'概念展开\'!$A$2:$A${last_flat},B{i},\'概念展开\'!$B$2:$B${last_flat},C{i},\'概念展开\'!$C$2:$C${last_flat},"核心")').font = Font(color=GREEN)
    ws.cell(i, 5, f'=COUNTIFS(\'概念展开\'!$A$2:$A${last_flat},B{i},\'概念展开\'!$B$2:$B${last_flat},C{i},\'概念展开\'!$C$2:$C${last_flat},"相关")').font = Font(color=GREEN)
    ws.cell(i, 6, f"=SUM(D{i}:E{i})")
    for j in range(2, 7):
        ws.cell(i, j).alignment = Alignment(horizontal="left" if j <= 3 else "center")
        if i % 2 == 0:
            ws.cell(i, j).fill = ALT_FILL
last_stats = 4 + len(concepts)
ws.cell(last_stats + 1, 2, "合计（概念次数，非股票去重数）").font = Font(bold=True)
for j in range(4, 7):
    col = get_column_letter(j)
    ws.cell(last_stats + 1, j, f"=SUM({col}5:{col}{last_stats})").font = Font(bold=True)
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 32
for col in ["D", "E", "F"]:
    ws.column_dimensions[col].width = 18
ws.freeze_panes = "B5"
ws.auto_filter.ref = f"B4:F{last_stats}"
save_check(ws, last_stats + 1)

ws = wb.create_sheet("封面", 0)
setup(ws, "机器人与物理AI概念深度清洗报告", "复核日期：2026-10-27｜范围：137只｜口径：产品直接性、商业化证据、主营相关度、分类边界", 7)
for col, width in zip(range(2, 9), [30, 18, 18, 18, 18, 18, 18]):
    ws.column_dimensions[get_column_letter(col)].width = width
metrics = [
    ("复核股票（只）", f"=COUNTA('清洗总表'!$B$5:$B${4 + len(audit)})"),
    ("保留机器人（只）", f"=COUNTA('机器人'!$B$5:$B${4 + len(robot)})"),
    ("归入物理AI（只）", f"=COUNTA('物理AI'!$B$5:$B${4 + len(physical)})"),
    ("完全摘除（只）", f"=COUNTA('摘除清单'!$B$5:$B${4 + len(removed)})"),
    ("迁移/新增物理AI动作（只）", f'=COUNTIF(\'物理AI\'!$E$5:$E${4 + len(physical)},"迁移")+COUNTIF(\'物理AI\'!$E$5:$E${4 + len(physical)},"新增")'),
    ("纠错动作（只）", f'=COUNTIF(\'清洗总表\'!$E$5:$E${4 + len(audit)},"纠错")'),
]
for i, (label, formula) in enumerate(metrics, 6):
    ws.cell(i, 2, label).font = Font(name="Microsoft YaHei", bold=True, color="333333")
    ws.cell(i, 3, formula).font = Font(name="Microsoft YaHei", size=14, bold=True, color=GREEN)
    ws.cell(i, 3).alignment = Alignment(horizontal="center")
    ws.row_dimensions[i].height = 26
ws.cell(14, 2, "工作表索引").font = Font(size=13, bold=True, color="122B49")
index_rows = [
    ("清洗总表", "137只旧分类、新分类、动作、判词、详细理由及证据"),
    ("机器人", f"{len(robot)}只保留的机器人本体、集成及硬件零部件公司"),
    ("物理AI", f"{len(physical)}只感知、控制、训练、数据、终端及具身本体公司"),
    ("摘除清单", f"{len(removed)}只错标、间接映射或研发预期退出公司"),
    ("概念统计", "各二级概念核心/相关数量；底层为精确展开数据"),
]
for i, (name, desc) in enumerate(index_rows, 15):
    ws.cell(i, 2, name).font = Font(bold=True, color=BLUE)
    ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=7)
    ws.cell(i, 3, desc).font = Font(color=BLUE)
notes = [
    "核心：相关产品是明确主营或主要收入应用；相关：有直接产品、订单、交付或收入，但仍属新业务或收入占比较低。",
    "摘除：仅参股、通用能力映射、无直接证据，或只有研发送样且无商业化；小额真实收入与纯研发预期分开处理。",
    "边界：硬件零部件留在机器人；具身本体、感知、数据、仿真、动作捕捉、智能控制、终端和通信模组归物理AI。",
    "主营资料：iFinD ifind_get_stock_info（2026-10-27）；逐股证据名称、日期与URL见各明细表Source列。",
]
for i, text in enumerate(notes, 22):
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=8)
    ws.cell(i, 2, text).font = Font(size=10, color="666666")
save_check(ws, 25)

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUT)
print(OUT)
print("sheets", wb.sheetnames)
print("rows", len(audit), len(robot), len(physical), len(removed), len(concepts), last_flat - 1)
