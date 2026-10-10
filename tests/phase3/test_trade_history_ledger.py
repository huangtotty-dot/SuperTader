# -*- coding: utf-8 -*-
"""历史成交解析 + 逐股盈亏·做T 台账 `core/trade_history.py` 的离线单测（2026-10-10，owner 需求3）。

铁律：全离线、纯文本喂入，不读任何真实文件、不联网。

运行：python tests/phase3/test_trade_history_ledger.py
"""
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

import core.trade_history as j  # noqa: E402

HEAD_L = "成交日期 证券代码 证券名称 操作 成交数量 成交均价 成交金额 可用余额 发生金额"
HEAD_R = "印花税 其他杂费 资金余额 合同编号 佣金 过户费 结算费"

# 交割单跨日降序、日内按成交先后 —— 这里照抄真实排版
MAIN = [
    HEAD_L,
    "20260215 1 测试甲 买入 0 0 5 50 -5",        # 纯费用行（qty=0）→ misc
    "20260210 1 测试甲 卖出 50 25 1250 50 1245",  # 部分卖出
    "20260201 1 测试甲 买入 100 20 2000 100 -2005",  # 清仓后再建
    "20260110 1 测试甲 卖出 300 12 3600 0 3590",  # 清仓
    "20260105 1 测试甲 买入 200 11 2200 300 -2205",
    "20260102 1 测试甲 买入 100 10 1000 100 -1005",
]

JOIN = [
    HEAD_L,
    "20260303 600000 测试丙 买入 100 10 1000 200 -1005.05",
    "20260302 600000 测试丙 卖出 100 11 1100 100 1093.85",
    "20260301 600000 测试丙 买入 200 9 1800 200 -1805.06",
    HEAD_R,
    "0 0 50000 11 5 0.05 0",
    "1.1 0 51000 12 5 0.05 0",
    "0 0 49000 13 5 0.06 0",
]

EDGE = [
    HEAD_L,
    "20260402 600001 测试丁 买入 100 10 1000 200 -1005",   # 余额对不上 → 预警
    "20260401 131801 Ｒ-007 卖出 40 1.49 4000 4000 -4000.2",  # 逆回购 → 剔除出盈亏
    HEAD_R,
    "0 0 1000 21 5 0.01 0",
    "0 0 5000 22 5 0 0",
]


class TestParseRows(unittest.TestCase):
    def test_01_代码补零与升序(self):
        tr = j.parse_rows(MAIN)
        self.assertEqual(len(tr), 6, "表头不应被当日志行")
        self.assertEqual([t["code"] for t in tr], ["000001"] * 6)
        self.assertEqual(tr[0]["date"], "2026-01-02")
        self.assertEqual(tr[-1]["date"], "2026-02-15", "应按日期升序重排")

    def test_02_字段解析(self):
        t = j.parse_rows(MAIN)[1]          # 2026-01-05 买入 200
        self.assertEqual(t["op"], "买入")
        self.assertEqual(t["qty"], 200)
        self.assertEqual(t["price"], 11.0)
        self.assertEqual(t["amount"], 2200.0)
        self.assertEqual(t["avail"], 300)
        self.assertEqual(t["net"], -2205.0)
        self.assertIsNone(t["fees"], "无右表时 fees 应为 None")

    def test_03_左右表按下标join(self):
        tr = j.parse_rows(JOIN)
        self.assertEqual(len(tr), 3)
        first = next(t for t in tr if t["date"] == "2026-03-03")   # 报价单首行
        self.assertIsNotNone(first["fees"])
        self.assertEqual(first["fee_total"], 5.05)
        self.assertEqual(first["cash_balance"], 50000.0)
        self.assertEqual(first["contract"], "11")

    def test_04_页标记与页眉被忽略(self):
        tr = j.parse_rows(["===== PAGE 3 ====="] + MAIN + ["===== PAGE 4 ====="])
        self.assertEqual(len(tr), 6)

    def test_05_垃圾行不被误判(self):
        tr = j.parse_rows([HEAD_L, "这是一行说明文字 不是成交 记录", "2026-01-02 乱码"])
        self.assertEqual(tr, [])


class TestBuildLedger(unittest.TestCase):
    def setUp(self):
        self.led = j.build_ledger(j.parse_rows(MAIN))
        self.st = self.led["stocks"][0]

    def test_01_已实现_建加清(self):
        # 买 1005 + 买 2205 = 3210，卖 3590 → +380
        self.assertAlmostEqual(self.st["realized"], 622.5, places=2)
        self.assertAlmostEqual(self.st["buy_amt"], 5215.0, places=2)   # 1005+2205+2005
        self.assertAlmostEqual(self.st["sell_amt"], 4835.0, places=2)  # 3590+1245

    def test_02_未平仓位与成本(self):
        # 清仓后 2026-02-01 再买 100@20（费5）→ 成本 2005；再卖 50 → 留 50 股 1002.5
        self.assertEqual(self.st["open_qty"], 50)
        self.assertAlmostEqual(self.st["cost_basis"], 1002.5, places=2)
        self.assertAlmostEqual(self.st["realized"], 380.0 + 242.5, places=2)

    def test_03_qty0杂项行进misc(self):
        self.assertAlmostEqual(self.st["misc"], -5.0, places=2)
        self.assertEqual(self.st["n_trades"], 6)

    def test_04_月度与累计(self):
        mon = {m["month"]: m["realized"] for m in self.led["monthly"]}
        self.assertAlmostEqual(mon["2026-01"], 380.0, places=2)
        self.assertAlmostEqual(mon["2026-02"], 242.5, places=2)
        cum = self.led["cum"]
        self.assertEqual(cum[0]["date"], "2026-01-10")
        self.assertAlmostEqual(cum[-1]["cum_realized"], 622.5, places=2)

    def test_04b_月度笔数含买入腿(self):
        by = {m["month"]: m for m in self.led["monthly"]}
        self.assertEqual(by["2026-01"]["n_trades"], 3, "01-02/01-05 买入 + 01-10 卖出")
        self.assertEqual(by["2026-02"]["n_trades"], 3)
        self.assertEqual(self.led["total_fees"], 0.0, "无右表 ⇒ 无费用可计")

    def test_09_交易总费用与月度费用(self):
        led = j.build_ledger(j.parse_rows(JOIN))
        # 右表三行费用：5.05 / 6.15 / 5.06
        self.assertAlmostEqual(led["total_fees"], 16.26, places=2)
        bd = led["fee_breakdown"]
        self.assertAlmostEqual(bd["commission"], 15.0, places=2)
        self.assertAlmostEqual(bd["stamp"], 1.1, places=2)
        self.assertAlmostEqual(bd["transfer"], 0.16, places=2)
        self.assertEqual(bd["settle"], 0.0)
        m = led["monthly"][0]
        self.assertEqual(m["month"], "2026-03")
        self.assertAlmostEqual(m["fees"], 16.26, places=2)
        self.assertEqual(m["n_trades"], 3)
        # 月度费用之和 == 总费用
        self.assertAlmostEqual(sum(x["fees"] for x in led["monthly"]),
                               led["total_fees"], places=2)

    def test_10_总费用覆盖逆回购(self):
        # 逆回购虽被剔除出逐股盈亏，其手续费仍计入「交易总费用」（真实现金支出）
        led = j.build_ledger(j.parse_rows(EDGE))
        self.assertAlmostEqual(led["total_fees"], 10.01, places=2)
        self.assertAlmostEqual(sum(s["fees"] for s in led["stocks"]), 5.01, places=2,
                               msg="逐股费用不含被剔除的逆回购")
        self.assertAlmostEqual(led["monthly"][0]["fees"], 10.01, places=2)

    def test_05_区间与无预警(self):
        r = self.led["range"]
        self.assertEqual(r["n_trades"], 6)
        self.assertEqual(r["n_stocks"], 1)
        self.assertEqual(self.led["parse_warnings"], [])

    def test_06_左右表join后已实现(self):
        led = j.build_ledger(j.parse_rows(JOIN))
        s = led["stocks"][0]
        # 买 1805.06/200 = 9.0253；卖 100 → 1093.85 − 902.53 = 191.32
        self.assertAlmostEqual(s["realized"], 191.32, places=2)
        self.assertEqual(s["open_qty"], 200)
        self.assertAlmostEqual(s["cost_basis"], 1907.58, places=2)
        self.assertAlmostEqual(s["fees"], 16.26, places=2)
        self.assertEqual(led["parse_warnings"], [])

    def test_07_逆回购剔除且余额不符报警(self):
        led = j.build_ledger(j.parse_rows(EDGE))
        self.assertEqual([s["code"] for s in led["stocks"]], ["600001"])
        self.assertEqual([e["code"] for e in led["excluded"]], ["131801"])
        self.assertAlmostEqual(led["total_excluded_net"], -4000.2, places=2)
        self.assertEqual(len(led["parse_warnings"]), 1)
        self.assertIn("可用余额", led["parse_warnings"][0])

    def test_08_空输入(self):
        led = j.build_ledger([])
        self.assertEqual(led["stocks"], [])
        self.assertEqual(led["total_realized"], 0.0)
        self.assertEqual(led["range"]["n_trades"], 0)


class TestTableRows(unittest.TestCase):
    """xls/xlsx 路径：表头驱动（券商列序/列名不一致）+ 多文件合并（2026-10-10 owner 追加）。"""

    def test_01_列序打乱也能认(self):
        rows = [
            ["历史交割单", "", "", "", "", ""],
            ["证券代码", "成交日期", "证券名称", "操作", "成交数量", "成交均价",
             "成交金额", "可用余额", "发生金额", "资金余额"],
            ["600000", "2026-01-02", "测试丙", "买入", 100, 10, 1000, 100, -1005, 49000],
            ["", "", "", "", "", "", "", "", "", ""],          # 空行
            ["600000", "2026-01-10", "测试丙", "卖出", 100, 12, 1200, 0, 1192, 50192],
        ]
        tr = j.parse_table_rows(rows)
        self.assertEqual(len(tr), 2, "标题行/空行不能当成交")
        self.assertEqual(tr[0]["code"], "600000")
        self.assertEqual(tr[0]["date"], "2026-01-02")
        self.assertEqual(tr[0]["net"], -1005.0)
        self.assertEqual(tr[0]["avail"], 100)
        self.assertEqual(tr[0]["cash_balance"], 49000.0)

    def test_02_同行费用列(self):
        rows = [
            ["成交日期", "证券代码", "证券名称", "操作", "成交数量", "成交均价",
             "成交金额", "可用余额", "发生金额", "印花税", "佣金", "过户费"],
            ["2026-02-10", "2176", "江特电机", "卖出", 1300, 9.0, 11700, 0, 11682, 11.7, 5, 1.3],
        ]
        t = j.parse_table_rows(rows)[0]
        self.assertEqual(t["code"], "002176")
        self.assertAlmostEqual(t["fee_total"], 18.0, places=2)
        self.assertEqual(t["fees"]["stamp"], 11.7)

    def test_03_无成交表头返回空(self):
        self.assertEqual(j.parse_table_rows([["随便", "写点", "东西"]]), [])

    def test_04_日期归一(self):
        f = j._date
        self.assertEqual(f("20260201"), "2026-02-01")
        self.assertEqual(f("2026-2-1"), "2026-02-01")
        self.assertEqual(f("2026/02/01"), "2026-02-01")
        self.assertEqual(f("2026-02-01 00:00:00"), "2026-02-01")

    def test_05_多份合并去重(self):
        a = j.parse_table_rows([
            ["成交日期", "证券代码", "证券名称", "操作", "成交数量", "成交均价", "成交金额", "可用余额", "发生金额"],
            ["2026-03-05", "600000", "甲", "卖出", 100, 11, 1100, 0, 1095],
            ["2026-03-01", "600000", "甲", "买入", 100, 10, 1000, 100, -1005],
        ])
        b = j.parse_table_rows([                # 与 a 完全重叠的第二期
            ["成交日期", "证券代码", "证券名称", "操作", "成交数量", "成交均价", "成交金额", "可用余额", "发生金额"],
            ["2026-03-01", "600000", "甲", "买入", 100, 10, 1000, 100, -1005],
            ["2026-03-08", "600000", "甲", "买入", 200, 9, 1800, 200, -1805],
        ])
        merged, dup = j.merge_trades([a, b])
        self.assertEqual(dup, 1, "重叠的那一笔应被丢掉")
        self.assertEqual([t["date"] for t in merged],
                         ["2026-03-01", "2026-03-05", "2026-03-08"])
        led = j.build_ledger(merged)
        self.assertEqual(led["range"]["n_trades"], 3)
        self.assertEqual(led["stocks"][0]["open_qty"], 200)


class TestFileReaders(unittest.TestCase):
    """.xlsx / 伪-.xls(HTML) 的真实读盘（openpyxl / bs4 都在 requirements 里）。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_01_xlsx读取(self):
        import openpyxl
        p = self.d / "a.xlsx"
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["证券代码", "成交日期", "证券名称", "操作", "成交数量", "成交均价",
                   "成交金额", "可用余额", "发生金额"])
        ws.append(["600000", "2026-01-02", "测试丙", "买入", 100, 10, 1000, 100, -1005])
        ws.append(["600000", "2026-01-10", "测试丙", "卖出", 100, 12, 1200, 0, 1192])
        wb.save(p)
        led = j.parse_excel(p)
        self.assertEqual(led["range"]["n_trades"], 2)
        self.assertAlmostEqual(led["total_realized"], 187.0, places=2)

    def test_02_伪xls是HTML(self):
        # 国内券商常见：扩展名 .xls、内容却是 <table> 拼的 HTML，且多用 GBK
        html = ("<html><body><table>"
                "<tr><th>成交日期</th><th>证券代码</th><th>证券名称</th><th>操作</th>"
                "<th>成交数量</th><th>成交均价</th><th>成交金额</th><th>可用余额</th><th>发生金额</th></tr>"
                "<tr><td>20260102</td><td>600000</td><td>测试丙</td><td>买入</td>"
                "<td>100</td><td>10</td><td>1000</td><td>100</td><td>-1005</td></tr>"
                "</table></body></html>")
        p = self.d / "b.xls"
        p.write_bytes(html.encode("gbk"))
        rows = j.read_excel_rows(p)
        self.assertTrue(any("成交日期" in "".join(r) for r in rows))
        tr = j.parse_table_rows(rows)
        self.assertEqual(len(tr), 1)
        self.assertEqual(tr[0]["date"], "2026-01-02")
        self.assertEqual(tr[0]["net"], -1005.0)

    def _xlsx(self, name, header, rows):
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(header)
        for r in rows:
            ws.append(list(r))
        wb.save(self.d / name)
        return self.d / name

    def test_03_带账户列的多份合并成一本(self):
        # 文件里**有资金账号列** ⇒ 同账户的多份文件合成一本账
        head = ["资金账号", "成交日期", "证券代码", "证券名称", "操作", "成交数量",
                "成交均价", "成交金额", "发生金额"]
        p1 = self._xlsx("p1.xlsx", head, [("A1", "2026-01-02", "600000", "测试丙", "买入", 100, 10, 1000, -1005)])
        p2 = self._xlsx("p2.xlsx", head, [("A1", "2026-01-10", "600000", "测试丙", "卖出", 100, 12, 1200, 1192)])
        led = j.build_from_files([p1, p2])
        self.assertEqual(led["sources"], ["p1.xlsx", "p2.xlsx"])
        self.assertEqual([a["account"] for a in led["accounts"]], ["A1"])
        self.assertEqual(led["range"]["n_trades"], 2)
        self.assertAlmostEqual(led["total_realized"], 187.0, places=2)
        self.assertEqual(led["stocks"][0]["open_qty"], 0)

    def test_03b_无账户列则一个文件当一个账户(self):
        # 没有资金账号列时，**文件名即账户名** —— 两个文件 = 两个账户，各自独立建账
        head = ["成交日期", "证券代码", "证券名称", "操作", "成交数量",
                "成交均价", "成交金额", "发生金额"]
        p1 = self._xlsx("东方.xlsx", head, [("2026-01-02", "600000", "测试丙", "买入", 100, 10, 1000, -1005)])
        p2 = self._xlsx("东莞.xlsx", head, [("2026-01-10", "600000", "测试丙", "卖出", 100, 12, 1200, 1192)])
        led = j.build_from_files([p1, p2])
        self.assertEqual(sorted(a["account"] for a in led["accounts"]), ["东方", "东莞"])
        # 东莞那笔卖出在东莞账本里是孤儿（没买过）⇒ 不能算进东方的已实现
        by = {s["account"]: s for s in led["stocks"]}
        self.assertEqual(by["东方"]["open_qty"], 100)
        self.assertEqual(by["东方"]["realized"], 0.0)

    def test_03c_同文件内完全相同的两笔不能去重(self):
        # 一张 20000 的单被撮成两笔 10000@同价 —— 去重会让股数凭空少一半
        head = ["成交日期", "证券代码", "证券名称", "操作", "成交数量",
                "成交均价", "成交金额", "发生金额"]
        row = ("2026-06-17", "588000", "科创50", "买入", 10000, 1.871, 18710, -18715)
        p = self._xlsx("a.xlsx", head, [row, row])
        led = j.build_from_files([p])
        self.assertEqual(led["range"]["n_trades"], 2)
        self.assertEqual(led["stocks"][0]["open_qty"], 20000)

    def test_04_坏文件不拖垮其余(self):
        p = self.d / "broken.xlsx"
        p.write_bytes(b"not an excel at all")
        led = j.build_from_files([p])
        self.assertEqual(led["stocks"], [])
        self.assertTrue(led["parse_warnings"])


class TestHistorySchema(unittest.TestCase):
    """历史成交（2026-10-10 owner 统一口径）：无成交金额/无发生金额/无余额列、B/S 方向、多账户。"""

    def _rows(self):
        return [
            ["资金账号", "成交日期", "成交时间", "证券代码", "证券名称", "买卖标志",
             "成交数量", "成交价格", "佣金"],
            ["A001", "2026-01-02", "09:31:05", "600000", "测试甲", "买入", 100, 10.0, 5],
            ["A001", "2026-01-02", "14:50:00", "600000", "测试甲", "卖出", 100, 10.5, 5],
            ["B002", "2026-01-02", "09:35:00", "600000", "测试甲", "B", 200, 11.0, 6],
            ["B002", "2026-02-03", "10:00:00", "600000", "测试甲", "S", 200, 12.0, 7],
        ]

    def test_01_缺列也能算(self):
        tr = j.parse_table_rows(self._rows())
        self.assertEqual(len(tr), 4)
        t = tr[0]
        self.assertEqual(t["account"], "A001")
        self.assertEqual(t["time"], "09:31:05")
        self.assertEqual(t["amount"], 1000.0, "无成交金额列时应按 价×量 补算")
        self.assertFalse(t["non_trade"], "补算后不能被误判成非资金过户")
        self.assertEqual(tr[2]["op"], "买入", "B 应识别为买入")
        self.assertEqual(tr[3]["op"], "卖出", "S 应识别为卖出")

    def test_02_多账户分开建账(self):
        led = j.build_ledger(j.parse_table_rows(self._rows()))
        self.assertEqual([a["account"] for a in led["accounts"]], ["B002", "A001"],
                         "按已实现降序（B002 是盈利的）")
        by = {s["account"]: s for s in led["stocks"]}
        self.assertEqual(set(by), {"A001", "B002"}, "两个账户的同名股票不能混成一个账本")
        self.assertAlmostEqual(by["A001"]["realized"], 40.0, places=2)   # (1050-5) - (1000+5)
        self.assertAlmostEqual(by["B002"]["realized"], 187.0, places=2)  # (2400-7)-(2200+6)
        self.assertEqual(by["A001"]["open_qty"], 0)
        self.assertEqual(by["B002"]["open_qty"], 0)

    def test_02b_账户收益率(self):
        led = j.build_ledger(j.parse_table_rows(self._rows()))
        by = {a["account"]: a for a in led["accounts"]}
        # 收益率分母 = 累计买入成本（A001 买入 100@10 + 费5 = 1005 ⇒ 40/1005）
        self.assertAlmostEqual(by["A001"]["buy_amt"], 1005.0, places=2)
        self.assertAlmostEqual(by["A001"]["pnl_pct"], 40.0 / 1005.0 * 100, places=2)
        self.assertAlmostEqual(by["B002"]["buy_amt"], 2206.0, places=2)
        self.assertAlmostEqual(by["B002"]["pnl_pct"], 187.0 / 2206.0 * 100, places=2)
        self.assertAlmostEqual(by["A001"]["open_cost"], 0.0, places=2)
        self.assertEqual(by["B002"]["n_stocks"], 1)

    def test_03_月度分录带账户(self):
        led = j.build_ledger(j.parse_table_rows(self._rows()))
        keys = [(m["account"], m["month"]) for m in led["monthly"]]
        self.assertIn(("A001", "2026-01"), keys)
        self.assertIn(("B002", "2026-02"), keys)
        self.assertIn("account", led["cum"][0])

    def test_04_做T收益_同日买卖配对(self):
        led = j.build_ledger(j.parse_table_rows(self._rows()))
        # A001 当天买 100@10(费5)、卖 100@10.5(费5) ⇒ 配 100 股，净价差 = 10.45 - 10.05 = 0.40
        self.assertAlmostEqual(led["total_t_pnl"], 40.0, places=2)
        self.assertEqual(len(led["t_monthly"]), 1)
        m = led["t_monthly"][0]
        self.assertEqual((m["account"], m["month"], m["n_pairs"]), ("A001", "2026-01", 1))
        self.assertEqual(led["t_stocks"][0]["code"], "600000")

    def test_05_非资金过户不进做T(self):
        rows = [
            ["成交日期", "证券代码", "证券名称", "操作", "成交数量", "成交均价", "成交金额"],
            ["2026-07-03", "588170", "科创半导", "买入", 9000, 3.65, 32850],
            ["2026-07-03", "588170", "科创半导", "卖出", 9000, 3.694, 33246],
            ["2026-07-03", "588170", "科创半导", "买入", 36000, 3.673, 0],   # 份额折算，非成交
        ]
        tr = j.parse_table_rows(rows)
        self.assertTrue(tr[-1]["non_trade"])
        led = j.build_ledger(tr)
        # 只按真实的 9000 股配对计价，折算行不参与
        self.assertAlmostEqual(led["total_t_pnl"], 9000 * (3.694 - 3.65), places=2)


class TestIncrementalImport(unittest.TestCase):
    """增量导入 + 「记忆已导入到哪一天」（2026-10-10 owner：导过就记住当天，下次导下一天）。

    券商是「每期一份」导出 ⇒ 日常只会导新那几天。若不做增量合并，只导新文件会把历史冲掉
    （实测踩过）。这里断言：导入是累加的、重导旧文件不改变结果、data_through 会推进。
    """

    def setUp(self):
        import t_gui
        self.t_gui = t_gui
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)
        self._orig = t_gui.TRADE_HISTORY_LEDGER
        t_gui.TRADE_HISTORY_LEDGER = self.d / "led.json"
        self.api = t_gui.Api()

    def tearDown(self):
        self.t_gui.TRADE_HISTORY_LEDGER = self._orig
        self._tmp.cleanup()

    def _mk(self, name, rows):
        import openpyxl
        p = self.d / name
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["资金账号", "成交日期", "证券代码", "证券名称", "买卖标志",
                   "成交数量", "成交价格", "佣金"])
        for r in rows:
            ws.append(list(r))
        wb.save(p)
        return p

    def test_01_增量累加且记忆日期(self):
        p1 = self._mk("d1.xlsx", [("A1", "2026-10-08", "600000", "甲", "买入", 100, 10.0, 5),
                                  ("A1", "2026-10-09", "600000", "甲", "卖出", 100, 11.0, 5)])
        r1 = self.api.import_trade_history(p1)
        self.assertTrue(r1["advanced"])
        self.assertEqual(r1["data_through"], "2026-10-09")
        self.assertIsNone(r1["prev_data_through"])
        self.assertNotIn("trades", r1, "原始成交行不下发前端")

        p2 = self._mk("d2.xlsx", [("A1", "2026-10-12", "600000", "甲", "买入", 200, 9.0, 5)])
        r2 = self.api.import_trade_history(p2)
        self.assertEqual(r2["range"]["start"], "2026-10-08", "旧成交必须还在")
        self.assertEqual(r2["range"]["n_trades"], 3)
        self.assertEqual(r2["data_through"], "2026-10-12")
        self.assertEqual(r2["prev_data_through"], "2026-10-09")
        self.assertTrue(r2["advanced"])

    def test_02_重导旧文件不重复计(self):
        p1 = self._mk("d1.xlsx", [("A1", "2026-10-08", "600000", "甲", "买入", 100, 10.0, 5),
                                  ("A1", "2026-10-09", "600000", "甲", "卖出", 100, 11.0, 5)])
        r1 = self.api.import_trade_history(p1)
        r2 = self.api.import_trade_history(p1)          # 同一份再导一次
        self.assertFalse(r2["advanced"], "没有新交易日应标 False")
        self.assertEqual(r2["range"]["n_trades"], r1["range"]["n_trades"])
        self.assertAlmostEqual(r2["total_realized"], r1["total_realized"], places=2)
        self.assertEqual(r2["dup_dropped"], 2)

    def test_03_账户列表来自历史成交(self):
        p1 = self._mk("d1.xlsx", [("A1", "2026-10-08", "600000", "甲", "买入", 100, 10.0, 5)])
        self.api.import_trade_history(p1)
        got = self.api.load_accounts()
        self.assertTrue(got["imported"])
        self.assertEqual([a["account"] for a in got["accounts"]], ["A1"])
        self.assertIn("pnl_pct", got["accounts"][0])


class TestRealBrokerFormats(unittest.TestCase):
    """真券商「xls」的五种真身 + 两处实测踩过的坑（2026-10-10，owner 提供真实截图后）。

    扩展名一律不可信，所以逐个 reader 试；这里各造一份最小样本验读得到。
    另有两个必须守住的细节：
      · 有的券商把**方向写进成交数量**（卖出 = -100）⇒ 数量必须取绝对值，否则卖出全被丢。
      · 同一张表里同时有「余额」「资金余额」「后证券余额」⇒ 列认领必须**别名长的优先**。
    """

    HEAD = ["交收时间", "成交日期", "流水号", "证券代码", "余额", "证券名称", "操作", "发生金额",
            "资金余额", "摘要", "后证券余额", "成交数量", "成交均价", "成交金额", "交易费用",
            "印花税", "其他费用", "成交编号", "币种", "佣金", "过户费", "余额"]
    ROWS = [
        ["20261009", "20261009", "3.2E+09", "600276", "0", "恒瑞医药", "卖出", "4591.65",
         "104675.5", "证券卖出", "0", "-100", "45.99", "4599", "7.35", "2.3", "0",
         "5E+10", "人民币", "5", "0.05", "104675"],
        ["20261009", "20261009", "3.2E+09", "2176", "1300", "江特电机", "买入", "-11055",
         "93620.47", "证券买入", "1300", "1300", "8.5", "11050", "5", "0", "0",
         "5E+10", "人民币", "5", "0", "93620"],
    ]

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _sml(self, name="a.xls", enc="gbk"):
        def esc(x):
            return str(x).replace("&", "&amp;").replace("<", "&lt;")
        body = "".join(
            "<Row>" + "".join(f'<Cell><Data ss:Type="String">{esc(c)}</Data></Cell>' for c in r) + "</Row>"
            for r in [self.HEAD] + self.ROWS)
        p = self.d / name
        p.write_bytes((
            '<?xml version="1.0"?><?mso-application progid="Excel.Sheet"?>'
            '<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" '
            'xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">'
            '<Worksheet ss:Name="Sheet1"><Table>' + body + '</Table></Worksheet></Workbook>').encode(enc))
        return p

    def test_01_excel2003xml真身(self):
        p = self._sml()
        self.assertEqual(j.sniff_format(p), "Excel 2003 XML(SpreadsheetML)")
        tr = j.trades_of(p)
        self.assertEqual(len(tr), 2, "卖出行不能被丢掉")

    def test_02_数量带符号_方向只认op列(self):
        tr = j.trades_of(self._sml())
        sell = next(t for t in tr if t["code"] == "600276")
        self.assertEqual(sell["op"], "卖出")
        self.assertEqual(sell["qty"], 100, "成交数量 -100 要取绝对值")
        self.assertEqual(sell["occur"], 4591.65, "发生金额为正 = 卖出回款")
        self.assertEqual(sell["price"], 45.99)

    def test_03_三列余额各归各位(self):
        tr = j.trades_of(self._sml())
        sell = next(t for t in tr if t["code"] == "600276")
        self.assertEqual(sell["avail"], 0, "后证券余额=0 → 持仓余额")
        self.assertEqual(sell["cash_balance"], 104675.5, "资金余额 ≠ 泛称「余额」")
        self.assertAlmostEqual(sell["fee_total"], 7.35, places=2)

    def test_04_html与tsv也认(self):
        head = ["成交日期", "证券代码", "证券名称", "操作", "成交数量", "成交均价", "成交金额"]
        rows = [["20261009", "600276", "恒瑞医药", "卖出", "-100", "45.99", "4599"]]
        p = self.d / "h.xls"
        p.write_bytes(("<html><table><tr>" + "".join(f"<th>{c}</th>" for c in head) + "</tr>"
                       + "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in rows)
                       + "</table></html>").encode("gbk"))
        self.assertEqual(j.sniff_format(p), "HTML 表格")
        self.assertEqual(len(j.trades_of(p)), 1)

        p2 = self.d / "t.xls"
        p2.write_bytes(("\n".join("\t".join(r) for r in [head] + rows)).encode("gbk"))
        self.assertEqual(len(j.trades_of(p2)), 1)

    def test_05_读不了时报出诊断(self):
        p = self.d / "bad.xls"
        p.write_bytes(b"\x00\x01\x02 not a spreadsheet at all")
        diag = []
        with self.assertRaises(Exception):
            j.trades_of(p, diag=diag)
        self.assertTrue(diag, "失败要留下可回传的诊断")
        self.assertIn("未知", diag[0])

    def _xlsx(self, name, header, rows):
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(header)
        for r in rows:
            ws.append(list(r))
        wb.save(self.d / name)
        return self.d / name

    def test_06_份额折算行不能记成本(self):
        """红股/份额折算行（成交金额=0）只加股数，**绝不能**按 价×量 记成本。

        实测：588170 一笔 36000 股@3.673 的折算行被记了 132,228 元假成本，
        该票已实现从 +4,118 变成 -128,110。
        """
        head = ["成交日期", "证券代码", "证券名称", "操作", "成交数量",
                "成交均价", "成交金额", "发生金额"]
        p = self._xlsx("z.xlsx", head, [
            ("2026-06-30", "588170", "科创半导", "买入", 1000, 3.60, 3600, -3605),
            ("2026-07-03", "588170", "科创半导", "买入", 2000, 3.60, 0, 0),      # 送股/折算
            ("2026-07-06", "588170", "科创半导", "卖出", 3000, 1.20, 3600, 3595),
        ])
        led = j.build_from_files([p])
        s = led["stocks"][0]
        self.assertEqual(s["open_qty"], 0)
        # 成本只有第一笔 3605；三笔 3000 股按 3605/3000 的成本卖出
        self.assertAlmostEqual(s["realized"], 3595.0 - 3605.0, places=2)

    def test_07_浮动盈亏与合计(self):
        p = self._xlsx("u.xlsx", ["成交日期", "证券代码", "证券名称", "操作",
                                  "成交数量", "成交均价", "成交金额", "发生金额"],
                       [("2026-01-02", "600000", "测试丙", "买入", 100, 10, 1000, -1005)])
        led = j.build_from_files([p])
        s = led["stocks"][0]
        self.assertEqual(s["open_qty"], 100)
        self.assertAlmostEqual(s["cost_basis"], 1005.0, places=2)
        # 给个固定价，浮动 = 12*100 - 1005 = 195
        led2 = j.build_ledger(led["trades"], price_fn=lambda c: 12.0)
        s2 = led2["stocks"][0]
        self.assertAlmostEqual(s2["unrealized"], 195.0, places=2)
        self.assertAlmostEqual(s2["total_pnl"], 195.0, places=2)
        self.assertAlmostEqual(led2["accounts"][0]["total_pnl"], 195.0, places=2)


if __name__ == "__main__":
    unittest.main(verbosity=2)
